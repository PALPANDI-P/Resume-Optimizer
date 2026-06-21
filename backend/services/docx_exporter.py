import docx
import re
from docx.shared import Pt, Inches, RGBColor


def create_docx(content, output_path, template_id="cc-001"):
    doc = docx.Document()

    TRACK_COLORS_DOCX = {
        'beg': [
            {'accent': '2563eb'}, {'accent': '3b82f6'}, {'accent': '16a34a'},
            {'accent': '64748b'}, {'accent': '1d4ed8'}, {'accent': '475569'},
            {'accent': '78716c'}, {'accent': '0891b2'}, {'accent': '65a30d'},
            {'accent': '7c3aed'},
        ],
        'exp': [
            {'accent': 'b45309'}, {'accent': 'b87333'}, {'accent': 'b91c1c'},
            {'accent': '78716c'}, {'accent': '1e40af'}, {'accent': 'a16207'},
            {'accent': '0369a1'}, {'accent': 'c9a227'}, {'accent': '6c757d'},
            {'accent': 'd97706'},
        ],
        'it': [
            {'accent': '0284c7'}, {'accent': '6d28d9'}, {'accent': '0d9488'},
            {'accent': '4f46e5'}, {'accent': '06b6d4'}, {'accent': '2563eb'},
            {'accent': '6366f1'}, {'accent': '0ea5e9'}, {'accent': '8b5cf6'},
            {'accent': '14b8a6'},
        ],
        'gov': [
            {'accent': '1e40af'}, {'accent': '15803d'}, {'accent': '57534e'},
            {'accent': '1d4ed8'}, {'accent': '047857'}, {'accent': '78716c'},
            {'accent': '0284c7'}, {'accent': '475569'}, {'accent': '4d7c0f'},
            {'accent': '65a30d'},
        ],
        'biz': [
            {'accent': '1e40af'}, {'accent': '0f766e'}, {'accent': '44403c'},
            {'accent': '1e40af'}, {'accent': '475569'}, {'accent': '78716c'},
            {'accent': '64748b'}, {'accent': '0284c7'}, {'accent': 'b91c1c'},
            {'accent': '2563eb'},
        ],
    }

    TRACK_FONTS_DOCX = {
        'beg': ['Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Times New Roman', 'Arial', 'Arial', 'Arial'],
        'exp': ['Times New Roman', 'Arial', 'Times New Roman', 'Arial', 'Arial', 'Times New Roman', 'Arial', 'Times New Roman', 'Arial', 'Times New Roman'],
        'it':  ['Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial'],
        'gov': ['Times New Roman', 'Times New Roman', 'Arial', 'Arial', 'Arial', 'Times New Roman', 'Arial', 'Arial', 'Times New Roman', 'Times New Roman'],
        'biz': ['Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Arial', 'Times New Roman', 'Arial'],
    }

    TRACK_SIZES = {
        'beg': [11, 11, 11, 11, 11, 10, 11, 11, 11, 10],
        'exp': [12, 11, 12, 11, 11, 12, 11, 12, 11, 12],
        'it':  [11, 11, 11, 11, 11, 11, 10, 10, 11, 11],
        'gov': [11, 11, 10, 10, 10, 11, 10, 10, 11, 11],
        'biz': [11, 11, 11, 11, 11, 10, 10, 10, 12, 11],
    }

    def build_docx_style(template_id):
        parts = template_id.split('-')
        if len(parts) != 3:
            return None
        track_key = parts[0]
        track_map = {'beg': 'beg', 'exp': 'exp', 'it': 'it', 'gov': 'gov', 'biz': 'biz'}
        key = track_map.get(track_key)
        if not key:
            return None
        try:
            idx = int(parts[2]) - 1
        except ValueError:
            return None
        mod = idx % 10
        if key not in TRACK_COLORS_DOCX or mod >= len(TRACK_COLORS_DOCX[key]):
            return None
        return {
            'font': TRACK_FONTS_DOCX[key][mod],
            'size': TRACK_SIZES[key][mod],
            'h_size': TRACK_SIZES[key][mod] + 3,
            'accent': TRACK_COLORS_DOCX[key][mod]['accent'],
        }

    STATIC_STYLES = {
        "ats-001": {"font": "Times New Roman", "size": 11, "h_size": 14, "accent": "1e40af"},
        "tech-001": {"font": "Arial", "size": 10, "h_size": 13, "accent": "2563eb"},
        "prof-001": {"font": "Times New Roman", "size": 11, "h_size": 14, "accent": "0f766e"},
        "ats-002": {"font": "Arial", "size": 10, "h_size": 13, "accent": "374151"},
        "beg-001": {"font": "Arial", "size": 11, "h_size": 14, "accent": "e11d48"},
        "exec-001": {"font": "Arial", "size": 12, "h_size": 15, "accent": "b87333"},
        "prof-002": {"font": "Times New Roman", "size": 11, "h_size": 14, "accent": "475569"},
        "tech-002": {"font": "Arial", "size": 10, "h_size": 13, "accent": "0284c7"},
        "exec-002": {"font": "Arial", "size": 11, "h_size": 14, "accent": "0f766e"},
        "beg-002": {"font": "Arial", "size": 11, "h_size": 14, "accent": "1e40af"},
        "tech-003": {"font": "Arial", "size": 10, "h_size": 13, "accent": "6d28d9"},
        "ats-003": {"font": "Arial", "size": 10, "h_size": 13, "accent": "1e40af"},
        "beg-003": {"font": "Arial", "size": 10, "h_size": 13, "accent": "16a34a"},
        "exec-003": {"font": "Times New Roman", "size": 12, "h_size": 15, "accent": "b45309"},
        "tech-004": {"font": "Arial", "size": 11, "h_size": 14, "accent": "d97706"},
    }

    def get_docx_style(tid):
        if tid in STATIC_STYLES:
            return STATIC_STYLES[tid]
        generated = build_docx_style(tid)
        if generated:
            return generated
        return {"font": "Arial", "size": 11, "h_size": 14, "accent": "000000"}

    config = get_docx_style(template_id)
    DEFAULT = {"font": "Arial", "size": 11, "h_size": 14, "accent": "000000"}

    # Improved margins: 0.6" top/bottom, 0.7" left/right
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = config["font"]
    style.font.size = Pt(config["size"])
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.15

    SECTION_HEADERS = {
        "PROFESSIONAL SUMMARY", "SUMMARY", "OBJECTIVE", "SKILLS", "WORK EXPERIENCE",
        "EXPERIENCE", "EDUCATION", "PROJECTS", "CERTIFICATIONS", "AWARDS",
        "PUBLICATIONS", "REFERENCES", "LANGUAGES", "INTERESTS", "VOLUNTEER",
        "PROFESSIONAL EXPERIENCE", "CAREER SUMMARY", "TECHNICAL SKILLS",
        "CORE COMPETENCIES", "EMPLOYMENT HISTORY",
    }

    ENTRY_SECTIONS = {
        "WORK EXPERIENCE", "EXPERIENCE", "EDUCATION", "PROJECTS",
        "CERTIFICATIONS", "PROFESSIONAL EXPERIENCE", "EMPLOYMENT HISTORY",
    }

    accent_color = RGBColor.from_string(config["accent"])

    lines = content.split("\n")
    current_section = None
    is_first_line = True
    body_started = False

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        normalized = line.upper().rstrip(":")
        is_header = normalized in SECTION_HEADERS or (
            line.isupper()
            and len(line) > 2
            and len(line) < 50
            and "@" not in line
            and "|" not in line
        )

        # Date-pattern detection for sub-headers
        date_match = re.search(r'((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{4}\s*[-–]\s*(?:Present|Current|\d{4}))|(\d{4}\s*[-–]\s*(?:Present|Current|\d{4}))', line)

        # First non-empty line = name
        if is_first_line and not is_header:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.bold = True
            run.font.size = Pt(config["h_size"] + 6)
            run.font.color.rgb = accent_color
            run.font.name = config["font"]
            p.space_after = Pt(4)
            p.space_before = Pt(0)
            is_first_line = False
            continue

        is_first_line = False

        if is_header:
            spacer = doc.add_paragraph()
            spacer.space_before = Pt(8)
            spacer.space_after = Pt(0)

            p = doc.add_paragraph()
            run = p.add_run(normalized)
            run.bold = True
            run.font.size = Pt(config["h_size"] + 1)
            run.font.color.rgb = accent_color
            run.font.name = config["font"]
            p.space_after = Pt(8)

            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement

            pBorder = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "2")
            bottom.set(qn("w:color"), config["accent"])
            pBorder.append(bottom)
            p.paragraph_format.element.get_or_add_pPr().append(pBorder)

            current_section = normalized
            continue

        # Bullet point
        if line.startswith("•") or line.startswith("-") or line.startswith("*"):
            p = doc.add_paragraph(style="List Bullet")
            content_text = line.lstrip("•-* ").strip()
            run = p.add_run(content_text)
            run.font.size = Pt(config["size"])
            run.font.name = config["font"]
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            continue

        # Sub-header detection (entry sections with " at ", "|", digits+hyphen, or date patterns)
        if current_section in ENTRY_SECTIONS:
            is_sub = False
            if " at " in line or "|" in line:
                is_sub = True
            elif date_match:
                is_sub = True
            elif any(ch.isdigit() for ch in line) and ("-" in line or "to" in line.lower()):
                is_sub = True

            if is_sub:
                p = doc.add_paragraph()
                run = p.add_run(line)
                run.bold = True
                run.font.size = Pt(config["size"] + 1)
                run.font.name = config["font"]
                p.space_after = Pt(3)
                p.space_before = Pt(2)
                continue

        # Contact info lines near top
        if not current_section and not body_started and ("@" in line or "phone" in line.lower() or "mobile" in line.lower() or "|" in line or any(loc in line.lower() for loc in ['location', 'address', 'city', 'state'])):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.size = Pt(config["size"] - 1)
            run.font.name = config["font"]
            run.font.color.rgb = RGBColor(100, 100, 100)
            p.alignment = 1  # center alignment
            p.space_after = Pt(2)
            p.space_before = Pt(0)
            continue

        body_started = True

        # Normal paragraph
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(0)

    doc.save(output_path)
